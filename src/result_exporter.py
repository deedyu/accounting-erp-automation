"""계산 값을 유지하며 검토용 Excel 결과를 생성한다."""
from pathlib import Path
from typing import Any
from datetime import date, datetime
import json
import math
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from src.daily_ledger import build_daily_balance_summary, build_monthly_balance_summary, build_side_by_side_ledger


def clean_value(value: Any) -> Any:
    """JSON·SQLite·Excel 저장용으로 결측과 날짜를 정규화한다."""
    if isinstance(value, dict):
        return {str(k): clean_value(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean_value(v) for v in value]
    if value is None or bool(pd.isna(value)):
        return None
    if isinstance(value, (date, datetime, pd.Timestamp)):
        return value.isoformat()
    if hasattr(value, "item"):
        return clean_value(value.item())
    return value


def json_text(value: Any) -> str:
    """비표준 NaN 없이 JSON 문자열을 생성한다."""
    return json.dumps(clean_value(value), ensure_ascii=False, allow_nan=False)


def write_excel(path: Path, sheets: dict[str, pd.DataFrame], *, totals: bool = False) -> None:
    """헤더 고정·필터·서식·오류 강조와 선택 합계행을 적용한다."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, frame in sheets.items():
        sheet = workbook.create_sheet(name[:31])
        columns = [str(c) for c in frame.columns]
        sheet.append(columns)
        money_cols = [i for i, c in enumerate(columns) if (("amount" in c and not c.endswith("count")) or c in {"debit_total", "credit_total", "difference"}) and not c.startswith("_")]
        for row_index, values in enumerate(frame.itertuples(index=False, name=None), start=2):
            record = []
            for value in values:
                clean = clean_value(value)
                record.append(json_text(clean) if isinstance(clean, (list, dict)) else clean)
            sheet.append(record)
            for cell in (sheet.cell(row_index, col) for col in range(1, len(columns) + 1)):
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
            status = values[columns.index("processing_status")] if "processing_status" in columns else ""
            if status in {"입력 불가", "검토 필요"}:
                fill = PatternFill("solid", fgColor="FCE4D6" if status == "입력 불가" else "FFF2CC")
                for cell in (sheet.cell(row_index, col) for col in range(1, len(columns) + 1)):
                    cell.fill = fill
        if totals and money_cols and name not in {"전체 합계", "정상 합계"}:
            total = [None] * len(columns)
            total[0] = "합계"
            for index in money_cols:
                total[index] = sum(v for v in frame.iloc[:, index] if isinstance(v, (int, float)) and not pd.isna(v))
            sheet.append(total)
            for cell in sheet[len(frame) + 2]:
                cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(max(1, len(columns)))}{max(1, len(frame) + 1)}"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="24476B")
        for index, column in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(45, max(14, len(column) + 3))
            for cells in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell = cells[0]
                if index - 1 in money_cols:
                    cell.number_format = '#,##0"원";[Red]-#,##0"원"'
                if column in {"transaction_date", "period_start", "period_end", "processed_at"} and cell.value and cell.value != "합계":
                    parsed = pd.to_datetime(cell.value, errors="coerce", format="mixed")
                    if not pd.isna(parsed):
                        cell.value = parsed.to_pydatetime()
                        cell.number_format = "yyyy-mm-dd"
    workbook.save(path)


def export_workbooks(result: dict[str, Any], accounts: pd.DataFrame) -> dict[str, Path]:
    """ERP·검토·오류·원본·일계표·요약 Excel을 함께 생성한다."""
    directory = Path(result["output_directory"])
    classified = result["classified_journal"]
    ready = classified.loc[classified.processing_status.eq("입력 가능")]
    ledger = build_side_by_side_ledger(classified, accounts)
    daily = build_daily_balance_summary(classified)
    ready_daily = build_daily_balance_summary(ready)
    monthly = build_monthly_balance_summary(daily, classified)
    ready_monthly = build_monthly_balance_summary(ready_daily, ready)
    result.update(daily_ledger=ledger, daily_summary=daily, ready_daily_summary=ready_daily,
                  monthly_summary=monthly, ready_monthly_summary=ready_monthly)
    paths = {}
    specifications = {
        "erp_excel": ("erp_upload.xlsx", {"ERP 분개": result["erp_upload"]}, True),
        "review_excel": ("review_queue.xlsx", {"검토 대기열": result["review_queue"]}, False),
        "errors_excel": ("validation_errors.xlsx", {"검증 오류": result["validation_result"]}, False),
        "ledger_excel": ("daily_ledger.xlsx", {"전체 분개장": ledger, "전체 일계표": daily, "정상 일계표": ready_daily,
                          "전체 합계": monthly, "정상 합계": ready_monthly}, True),
        "summary_excel": ("processing_summary.xlsx", {"처리 요약": result["summary"], "ERP 요약": result["erp_summary"]}, False),
        "source_excel": ("source_details.xlsx", {"원본 시트": result["standardization"]["raw_dataframe"],
                          "미매핑 값": result["standardization"]["unmapped_data"], "분리한 행": result["standardization"]["excluded_rows"]}, False),
    }
    for key, (name, sheets, totals) in specifications.items():
        path = directory / name
        write_excel(path, sheets, totals=totals)
        paths[key] = path
    return paths
